"""Merge per-game research JSON files (data/*.json) into games_analysis.xlsx and games_analysis.csv.

Usage: python3 build_spreadsheet.py
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"

# (json key, column header, column width)
COLUMNS = [
    ("group", "Group", 7),
    ("id", "ID", 6),
    ("title", "Title", 22),
    ("genre", "Genre / subgenre", 26),
    ("release_date", "Release date", 18),
    ("early_access_date", "Early access date", 14),
    ("window_note", "Time-window note", 22),
    ("developer", "Developer", 22),
    ("team_size", "Team size", 22),
    ("country", "Country", 12),
    ("developer_background", "Developer background", 40),
    ("publisher", "Publisher", 18),
    ("platforms_and_store_links", "Platforms & store links", 40),
    ("business_model", "Business model", 30),
    ("price_usd", "Price (USD)", 16),
    ("units_or_downloads", "Units sold / downloads / players", 45),
    ("peak_concurrent_players", "Peak concurrent players", 35),
    ("steam_reviews", "Steam reviews (or store ratings)", 30),
    ("implied_review_multiplier", "Implied units-per-review", 22),
    ("gross_revenue", "Gross revenue", 45),
    ("net_revenue", "Net revenue", 45),
    ("production_cost", "Production cost", 40),
    ("development_time", "Development time", 25),
    ("why_popular", "Why it became popular (hook / mechanic)", 55),
    ("marketing_channels", "Marketing channels & timing", 45),
    ("engine", "Engine", 18),
    ("art_style", "Art style", 22),
    ("ai_feasibility_score", "AI feasibility (1-5)", 10),
    ("ai_feasibility_justification", "AI feasibility justification & toolset", 60),
    ("solo_ai_mvp_cost_time", "Solo+AI MVP: cost & time", 30),
    ("solo_ai_polished_cost_time", "Solo+AI polished: cost & time", 30),
    ("data_conflicts", "Data conflicts / reconciliation", 45),
    ("notes", "Notes for a solo creator", 45),
    ("sources", "Sources", 70),
]

GROUP_FILL = {"A": "DCE6F1", "B": "E2EFDA"}


def load_games():
    games = []
    for path in sorted(DATA.glob("*.json")):
        with path.open(encoding="utf-8") as f:
            game = json.load(f)
        game["_file"] = path.name
        games.append(game)
    games.sort(key=lambda g: (g.get("group", "Z"), g.get("id", "")))
    return games


def cell_value(game, key):
    value = game.get(key, "no data")
    if key == "sources":
        return "\n".join(
            f"{s.get('title', '').strip()} — {s.get('url', '').strip()}" for s in value or []
        )
    if value is None or value == "":
        return "no data"
    return value


def write_csv(games):
    with (ROOT / "games_analysis.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([header for _, header, _ in COLUMNS])
        for game in games:
            writer.writerow([cell_value(game, key) for key, _, _ in COLUMNS])


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "D2"
    ws.row_dimensions[1].height = 42


def write_games_sheet(wb, games):
    ws = wb.active
    ws.title = "Games"
    ws.append([header for _, header, _ in COLUMNS])
    for game in games:
        ws.append([cell_value(game, key) for key, _, _ in COLUMNS])
    style_header(ws, len(COLUMNS))
    for idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        fill = PatternFill("solid", fgColor=GROUP_FILL.get(row[0].value, "FFFFFF"))
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[:3]:
            c.fill = fill
            c.font = Font(bold=True)
    ws.auto_filter.ref = ws.dimensions


def write_summary_sheet(wb, games):
    """Numeric, sortable best-estimate view (from summary_numbers.json, filled after cross-checking)."""
    path = ROOT / "summary_numbers.json"
    if not path.exists():
        return
    with path.open(encoding="utf-8") as f:
        numbers = json.load(f)
    cols = [
        ("id", "ID", 6),
        ("group", "Group", 7),
        ("title", "Title", 26),
        ("platform", "Main platform", 14),
        ("release", "Release", 12),
        ("team", "Team size (people)", 10),
        ("price_usd", "Launch list price (USD)", 10),
        ("units_m", "Units / players (millions, best est.)", 14),
        ("units_label", "Units label", 12),
        ("peak_ccu_k", "Peak CCU (thousands)", 12),
        ("gross_musd", "Gross revenue (US$ M, best est.)", 14),
        ("gross_label", "Gross label", 12),
        ("net_musd", "Net revenue (US$ M, best est.)", 14),
        ("dev_months", "Dev time (months)", 10),
        ("ai_score", "AI feasibility (1-5)", 10),
        ("mvp_weeks", "Solo+AI MVP (weeks)", 10),
        ("comment", "Comment", 60),
    ]
    ws = wb.create_sheet("Summary (numeric)")
    ws.append([h for _, h, _ in cols])
    by_id = {g["id"]: g for g in games}
    for gid in sorted(numbers, key=lambda k: (numbers[k].get("group", "Z"), k)):
        row = dict(numbers[gid])
        row["id"] = gid
        row.setdefault("title", by_id.get(gid, {}).get("title", ""))
        ws.append([row.get(k, "") for k, _, _ in cols])
    style_header(ws, len(cols))
    for idx, (_, _, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions


def write_sources_sheet(wb, games):
    ws = wb.create_sheet("Sources")
    ws.append(["ID", "Title", "Source", "URL"])
    for game in games:
        for s in game.get("sources", []) or []:
            ws.append([game.get("id"), game.get("title"), s.get("title", ""), s.get("url", "")])
    style_header(ws, 4)
    for col, width in zip("ABCD", (6, 24, 60, 90)):
        ws.column_dimensions[col].width = width


def write_methodology_sheet(wb):
    ws = wb.create_sheet("Methodology")
    ws.column_dimensions["A"].width = 140
    text = (ROOT / "research_spec.md").read_text(encoding="utf-8")
    start = text.find("## Labelling rules")
    end = text.find("## Output:")
    lines = ["Research date: 2026-10-08. Window: ~April 2025 – October 2026.", ""]
    lines += text[start:end].strip().splitlines()
    for line in lines:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)


def main():
    games = load_games()
    if not games:
        raise SystemExit("no data/*.json files found")
    write_csv(games)
    wb = Workbook()
    write_games_sheet(wb, games)
    write_summary_sheet(wb, games)
    write_sources_sheet(wb, games)
    write_methodology_sheet(wb)
    wb.save(ROOT / "games_analysis.xlsx")
    groups = {}
    for g in games:
        groups[g.get("group")] = groups.get(g.get("group"), 0) + 1
    print(f"wrote {len(games)} games {groups} -> games_analysis.xlsx / games_analysis.csv")


if __name__ == "__main__":
    main()
